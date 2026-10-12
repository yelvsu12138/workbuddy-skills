# -*- coding: utf-8 -*-
"""
实验报告填写脚本：把内容写入西南大学实验报告模板主表。

用法：
    python3 fill_report.py --template <模板.docx> --content <内容.json> --output <输出.docx>

内容 JSON 结构：
{
  "project_name": "实验项目名称",
  "experiment_type": ["验证性"],              // 可含 验证性/设计性/综合性，勾选对应项
  "purpose": ["实验目的段落1", "段落2"],
  "process": [["加粗小标题", true], ["正文要点", false]],
  "analysis": [["加粗小标题", true], ["正文要点", false]]
}
"""
import argparse
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

TYPES = ['验证性', '设计性', '综合性']


def set_run(run, size=10.5, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')


def fill_name(t1, name):
    for row in t1.rows:
        if '实验项目名称' in row.cells[0].text:
            cell = row.cells[2]
            cell.text = ''
            p = cell.paragraphs[0]
            p.text = ''
            run = p.add_run(name)
            set_run(run, 12, True)
            return True
    return False


def fill_type(t1, selected):
    for row in t1.rows:
        if any('实验类型' in c.text for c in row.cells):
            cell = row.cells[4]
            parts = []
            for t in TYPES:
                mark = '☑' if t in selected else '☐'
                parts.append(' %s %s ' % (mark, t))
            cell.text = ''
            p = cell.paragraphs[0]
            p.text = ''
            run = p.add_run(''.join(parts))
            set_run(run, 10.5, False)
            return True
    return False


def add_blocks(cell, blocks):
    for item in blocks:
        text, bold = (item if isinstance(item, (list, tuple)) else (item, False))
        p = cell.add_paragraph()
        run = p.add_run(text)
        set_run(run, 10.5, bold)


def fill_section(t1, label, blocks):
    for row in t1.rows:
        if row.cells[0].text.strip() == label:
            add_blocks(row.cells[0], blocks)
            return True
    return False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True)
    ap.add_argument('--content', required=True)
    ap.add_argument('--output', required=True)
    args = ap.parse_args()

    with open(args.content, encoding='utf-8') as f:
        c = json.load(f)

    doc = Document(args.template)
    t1 = doc.tables[0]

    assert fill_name(t1, c.get('project_name', '')), '未找到「实验项目名称」单元格'
    assert fill_type(t1, c.get('experiment_type', ['验证性'])), '未找到「实验类型」单元格'
    assert fill_section(t1, '实验目的', c.get('purpose', [])), '未找到「实验目的」'
    assert fill_section(t1, '实验内容及过程', c.get('process', [])), '未找到「实验内容及过程」'
    assert fill_section(t1, '实验分析及总结', c.get('analysis', [])), '未找到「实验分析及总结」'

    doc.save(args.output)
    print('OK: 已写入 %s' % args.output)


if __name__ == '__main__':
    main()
